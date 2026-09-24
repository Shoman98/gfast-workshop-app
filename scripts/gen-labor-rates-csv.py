"""
Generate labor rate CSVs from Excel source files.

The Excel files use cross-sheet formulas (Part_Hours_Reference, Rate_Card,
Make_Tier_Map) that openpyxl cannot evaluate. This script reconstructs the
values manually and writes flat CSVs ready for Supabase import.

Usage:
  python3 scripts/gen-labor-rates-csv.py            # generate both
  python3 scripts/gen-labor-rates-csv.py --type repair
  python3 scripts/gen-labor-rates-csv.py --type replace
"""

import argparse
import csv
import datetime
import openpyxl

CONFIGS = {
    'repair': {
        'input':      '/Users/User/Downloads/01Parts_Labor_Rates_Repair_by_Vehicle_DYNAMIC.xlsx',
        'output':     '/Users/User/Downloads/labor_rates_repair_import.csv',
        'sheet':      'Labor_Rates',
        'total_col':  'total_repair_hrs',
        'has_part_price': False,
        # Column indices in the main sheet (0-based)
        'last_updated_idx':  16,
        'vehicle_make_idx':  17,
        'vehicle_model_idx': 18,
        'vehicle_year_idx':  19,
    },
    'replace': {
        'input':      '/Users/User/Downloads/01Parts_Labor_Rates_Replace_by_Vehicle_DYNAMIC.xlsx',
        'output':     '/Users/User/Downloads/labor_rates_replace_import.csv',
        'sheet':      'Labor_Rates_Replace',
        'total_col':  'total_replace_hrs',
        'has_part_price': True,
        'part_price_idx': 16,
        'last_updated_idx':  17,
        'vehicle_make_idx':  18,
        'vehicle_model_idx': 19,
        'vehicle_year_idx':  20,
    },
}

HOUR_FIELDS = [
    'refitting_labor_hrs', 'dent_hrs', 'paint_hrs', 'elec_hrs', 'intr_hrs',
    'cooling_hrs', 'susp_hrs', 'mechanical_hrs', 'glass_hrs',
]


def build_lookups(wb):
    # Part hours: part_id -> {field: value}. Columns are mapped BY HEADER NAME,
    # never by fixed position: the Replace and Repair source sheets order the
    # labor columns differently (Replace goes refitting, elec, intr, ... paint,
    # dent), so positional reads silently scrambled the categories — e.g.
    # Replace's intr_hrs column was being written into paint_hrs, which is why
    # airbags showed up under "أعمال دهان".
    part_hours = {}
    ph_idx = None
    for i, row in enumerate(wb['Part_Hours_Reference'].iter_rows(values_only=True)):
        if i == 0:
            ph_idx = {h: k for k, h in enumerate(row) if h}
            missing = [f for f in HOUR_FIELDS if f not in ph_idx]
            if missing:
                raise SystemExit(f'Part_Hours_Reference missing columns: {missing}')
            continue
        if row[0]:
            part_hours[row[0]] = {f: row[ph_idx[f]] for f in HOUR_FIELDS}

    # Make -> tier
    make_tier = {}
    for row in wb['Make_Tier_Map'].iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:
            make_tier[row[0]] = row[1]

    # Tier -> hourly rates
    tier_rates = {}
    for row in wb['Rate_Card'].iter_rows(min_row=4, max_row=7, values_only=True):
        if row[0] and row[1]:
            tier_rates[row[0]] = {'labor': row[1], 'paint': row[2]}

    return part_hours, make_tier, tier_rates


def generate(cfg):
    print(f'\n📂 Reading {cfg["input"]}...')
    wb = openpyxl.load_workbook(cfg['input'], data_only=True)
    part_hours, make_tier, tier_rates = build_lookups(wb)
    print(f'   {len(part_hours)} parts, {len(make_tier)} makes, {list(tier_rates)} tiers')

    ws = wb[cfg['sheet']]
    written = 0

    base_cols = [
        'part_id', 'part_name_ar', 'part_name_en', 'category',
        'refitting_labor_hrs', 'dent_hrs', 'paint_hrs', 'elec_hrs',
        'intr_hrs', 'cooling_hrs', 'susp_hrs', 'mechanical_hrs', 'glass_hrs',
        cfg['total_col'], 'hr_price_egp', 'paint_1hr_price_egp',
    ]
    if cfg['has_part_price']:
        base_cols.append('part_price')
    base_cols += ['vehicle_make', 'vehicle_model', 'vehicle_year', 'last_updated']

    with open(cfg['output'], 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(base_cols)

        for row in ws.iter_rows(min_row=2, values_only=True):
            part_id       = row[0]
            part_name_ar  = row[1]
            part_name_en  = row[2]
            category      = row[3]
            last_updated  = row[cfg['last_updated_idx']]
            vehicle_make  = row[cfg['vehicle_make_idx']]
            vehicle_model = row[cfg['vehicle_model_idx']]
            vehicle_year  = row[cfg['vehicle_year_idx']]
            part_price    = row[cfg['part_price_idx']] if cfg['has_part_price'] else None

            if not part_id or not vehicle_make or not vehicle_model:
                continue

            hrs = part_hours.get(part_id, {})
            ref = hrs.get('refitting_labor_hrs') or None
            dnt = hrs.get('dent_hrs')            or None
            pnt = hrs.get('paint_hrs')           or None
            elc = hrs.get('elec_hrs')            or None
            itr = hrs.get('intr_hrs')            or None
            col = hrs.get('cooling_hrs')         or None
            ssp = hrs.get('susp_hrs')            or None
            mec = hrs.get('mechanical_hrs')      or None
            gls = hrs.get('glass_hrs')           or None
            total = sum(v for v in [ref, dnt, pnt, elc, itr, col, ssp, mec, gls] if v)

            tier        = make_tier.get(vehicle_make, 'Mid')
            tier_rate   = tier_rates.get(tier, tier_rates.get('Mid', {'labor': 350, 'paint': 2000}))
            hr_price    = tier_rate['labor']
            paint_1hr   = tier_rate.get('paint', 2000)

            if isinstance(last_updated, datetime.datetime):
                last_updated = last_updated.strftime('%Y-%m-%d')

            out = [
                part_id, part_name_ar, part_name_en, category,
                ref or '', dnt or '', pnt or '', elc or '',
                itr or '', col or '', ssp or '', mec or '', gls or '',
                total or '', hr_price, paint_1hr,
            ]
            if cfg['has_part_price']:
                out.append(part_price if part_price is not None else '')
            out += [vehicle_make, vehicle_model, vehicle_year or '2000-2027', last_updated or '']

            writer.writerow(out)
            written += 1

    import os
    size = os.path.getsize(cfg['output']) / 1024 / 1024
    print(f'   ✅ {written:,} rows → {cfg["output"]} ({size:.1f} MB)')


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--type', choices=['repair', 'replace'], help='Which table to generate (default: both)')
    args = ap.parse_args()

    targets = [args.type] if args.type else ['repair', 'replace']
    for t in targets:
        generate(CONFIGS[t])
    print('\n🎉 Done!')

if __name__ == '__main__':
    main()
